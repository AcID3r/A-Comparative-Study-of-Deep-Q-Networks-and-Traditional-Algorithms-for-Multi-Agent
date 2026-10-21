"""Independently compare Word table contents, source blocks, assets and layout."""
import argparse
import json
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from evaluation.final_protocol import digest
from evaluation.team_protocol import ROOT
from paper.build_word import REFERENCE,REFERENCE_HASH,roman


def normalize(text): return ' '.join(text.replace('\u200b','').split())


def check(source,provenance_path,output):
    source,provenance_path,output=map(Path,(source,provenance_path,output))
    provenance=json.loads(provenance_path.read_text(encoding='utf-8'))
    assert digest(source)==provenance['docx_sha256']
    assert digest(REFERENCE)==provenance['reference_sha256']==REFERENCE_HASH
    markdown=(ROOT/'paper/paper.md').read_text(encoding='utf-8')
    assert digest(ROOT/'paper/paper.md')==provenance['source_sha256']
    assert digest(ROOT/'results/analysis/CP-07b/report.json')==provenance['analysis_sha256']
    document=Document(source)
    paragraphs=[p for p in document.paragraphs if p.text.strip() and not p._p.findall('.//'+qn('m:oMath'))]
    expected=[b['text'] for b in provenance['blocks'] if b['kind']=='paragraph']
    assert [normalize(p.text) for p in paragraphs]==[normalize(t) for t in expected],'Source paragraphs changed/missing'
    tables=[]; lines=markdown.splitlines(); index=0
    while index<len(lines):
        if not lines[index].strip().startswith('|'): index+=1; continue
        rows=[]
        while index<len(lines) and lines[index].strip().startswith('|'):
            cells=[v.strip() for v in lines[index].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',v) for v in cells): rows.append(cells)
            index+=1
        tables.append(rows)
    assert len(tables)==len(document.tables)==10
    for table,rows in zip(document.tables,tables):
        actual=[[normalize(cell.text) for cell in row.cells] for row in table.rows]
        assert actual==[[normalize(value) for value in row] for row in rows],'Measured table changed'
        assert sum(col.width.twips for col in table.columns)==4869
        assert table.rows[0]._tr.find(qn('w:trPr')).find(qn('w:tblHeader')) is not None
    assert len(document.sections)==2
    for index,section in enumerate(document.sections):
        profile=provenance['formatting_profile']['sections'][index]
        assert section.page_width.twips==profile['width'] and section.page_height.twips==profile['height']
        for name,key in [('top_margin','top'),('bottom_margin','bottom'),('left_margin','left'),('right_margin','right')]:
            assert getattr(section,name).twips==profile[key]
        cols=section._sectPr.find(qn('w:cols'))
        assert int(cols.get(qn('w:num'),'1'))==index+1
        if index==1: assert cols.get(qn('w:space'))=='440'
    assert document.styles['Normal'].font.name=='Times New Roman'
    assert document.paragraphs[0].runs[0].font.size.pt==14
    text='\n'.join(p.text for p in paragraphs)
    for i,title in enumerate(('Introduction','Related Work','Problem Formulation','Methodology',
                              'Experimental Setup','Results','Discussion','Limitations','Conclusion'),1):
        assert roman(i)+'. '+title in text
    assert 'Madhav Joshi' in text and 'Marwadi University' in text
    assert not re.search(r'\bTODO\b|\bTBD\b|\[insert|\bQENN\b|\bqubit\b|\bNISQ\b',text,re.I)
    assert len(document.inline_shapes)==7
    with ZipFile(source) as package:
        xml=etree.fromstring(package.read('word/document.xml'))
        ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
            'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        equations=xml.findall('.//m:oMath',ns); assert len(equations)==2
        assert equations[0].find('m:sSub',ns) is not None
        assert equations[1].find('m:f',ns) is not None and equations[1].find('m:nary',ns) is not None
        assert [p.text for p in document.paragraphs if p._p.findall('.//'+qn('m:oMath'))]==['\t(1)','\t(2)']
        assert len(xml.findall('.//w:hyperlink',ns))==9
        expected_images=[b['sha256'] for b in provenance['blocks'] if b['kind']=='figure']
        import hashlib
        images=[hashlib.sha256(package.read(name)).hexdigest() for name in package.namelist() if name.startswith('word/media/')]
        assert sorted(images)==sorted(expected_images)
    report=json.loads((ROOT/'results/analysis/CP-07b/report.json').read_text())
    for method in ('astar','solo','shared'):
        group=next(g for g in report['groups']['overall'] if g['method']==method)
        assert f'{100*group["success_rate"]:.2f}%' in text
    native_path=source.parent/'word_native_render.json'
    native=json.loads(native_path.read_text()) if native_path.exists() else None
    if native is not None: assert native['source_sha256']==digest(source)
    result=dict(date='2026-10-07',docx_sha256=digest(source),
        reference_unchanged=True,source_paragraphs_checked=len(expected),tables_checked=10,
        figure_hashes_checked=7,editable_numbered_equations=2,reference_hyperlinks=9,
        headings_columns_geometry_match=True,all_measured_cells_match=True,
        confirmed_author=True,body_size_assumption=provenance['formatting_profile']['body_size_note'],
        native_render=native,status='Structural checks passed; native Word pagination review pending')
    output.write_text(json.dumps(result,indent=2)+'\n'); print(json.dumps(result,indent=2))
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'paper/Graph_Search_Research_Paper.docx')
    parser.add_argument('--provenance',type=Path,default=ROOT/'paper/word_provenance.json')
    parser.add_argument('--output',type=Path,default=ROOT/'results/validation/CP-09-checks.json')
    args=parser.parse_args(); check(args.source,args.provenance,args.output)
