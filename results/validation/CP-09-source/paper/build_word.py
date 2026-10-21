"""Build a separate reference-format DOCX from reviewed, audited CP-08 content."""
import argparse
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from evaluation.final_protocol import digest
from evaluation.team_protocol import ROOT
from paper.word_figures import generate

REFERENCE = ROOT/'QENN_Research_Paper_Madhav_Joshi (1).docx'
REFERENCE_HASH = '0eb7c34b6358da752b52c38fd9af9daead8469c2df2216df1c272f07c47392f8'
OMISSIONS = ['This manuscript is a completed scientific source and readable review PDF. ',
             'Venue-specific formatting and a separate reference-format Word deliverable remain the subsequent checkpoint.']


def roman(number):
    result = ''
    for value, symbol in ((10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')):
        while number>=value: result+=symbol; number-=value
    return result


def editorial(text):
    for sentence in OMISSIONS: text=text.replace(sentence,'')
    text=re.sub(r'Table (\d+)',lambda m:'Table '+roman(int(m[1])+1),text)
    text=re.sub(r'Figure (\d+)',r'Fig. \1',text)
    text=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1',text)
    return text.replace('**','').replace('`','').strip()


def profile():
    doc=Document(REFERENCE)
    sections=[]
    for section in doc.sections:
        cols=section._sectPr.find(qn('w:cols'))
        sections.append(dict(width=section.page_width.twips,height=section.page_height.twips,
            top=section.top_margin.twips,bottom=section.bottom_margin.twips,
            left=section.left_margin.twips,right=section.right_margin.twips,
            header=section.header_distance.twips,footer=section.footer_distance.twips,
            columns=int(cols.get(qn('w:num'),'1')),gap=int(cols.get(qn('w:space'),'720'))))
    return dict(reference_sha256=digest(REFERENCE),sections=sections,font='Times New Roman',
        title_points=doc.paragraphs[0].runs[0].font.size.pt,
        author_points=doc.paragraphs[1].runs[0].font.size.pt,
        affiliation_points=doc.paragraphs[3].runs[0].font.size.pt,
        caption_points=doc.paragraphs[13].runs[0].font.size.pt,
        table_points=doc.tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt,
        body_points=11,body_size_note='Reference omits explicit body size; 11pt Word default made explicit.',
        omitted_unconfirmed_author_fields=['department','postal location','email'],
        table_convention='Roman numbers, centered label/title',
        equation_convention='right-aligned parenthesized numbers',headers_footers='blank like reference')


def math_run(parent,text):
    run=OxmlElement('m:r'); value=OxmlElement('m:t'); value.text=text; run.append(value); parent.append(run)


def equation(paragraph,number):
    math=OxmlElement('m:oMath')
    if number==1:
        math_run(math,'y = r + γ(1 − d) ')
        sub=OxmlElement('m:sSub')
        for tag,value in [('m:e','max'),('m:sub','a′')]:
            item=OxmlElement(tag); math_run(item,value); sub.append(item)
        math.append(sub); math_run(math,' Q(o′, a′; θ⁻)')
    else:
        math_run(math,'ℒ(θ) = ')
        fraction=OxmlElement('m:f')
        for tag,value in [('m:num','1'),('m:den','B')]:
            item=OxmlElement(tag); math_run(item,value); fraction.append(item)
        math.append(fraction)
        summation=OxmlElement('m:nary'); props=OxmlElement('m:naryPr')
        for tag,value in [('m:chr','∑'),('m:limLoc','subSup')]:
            item=OxmlElement(tag); item.set(qn('m:val'),value); props.append(item)
        summation.append(props)
        for tag,value in [('m:sub','j=1'),('m:sup','B'),('m:e','[yⱼ − Q(oⱼ, aⱼ; θ)]²')]:
            item=OxmlElement(tag); math_run(item,value); summation.append(item)
        math.append(summation)
    for run in math.iter(qn('m:r')):
        props=OxmlElement('w:rPr'); size=OxmlElement('w:sz'); size.set(qn('w:val'),'19'); props.append(size)
        run.insert(0,props)
    paragraph._p.append(math)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Twips(4869),2)
    paragraph.add_run('\t'+f'({number})')


def build(source=ROOT/'paper/paper.md',output=ROOT/'paper/Graph_Search_Research_Paper.docx'):
    source,output=Path(source).resolve(),Path(output).resolve()
    if output==REFERENCE or digest(REFERENCE)!=REFERENCE_HASH: raise ValueError('Reference must remain unchanged')
    audit=json.loads((ROOT/'results/validation/CP-07b-checks.json').read_text())
    cp08=json.loads((ROOT/'paper/manuscript_provenance.json').read_text())
    assert audit['report_sha256']==digest(ROOT/'results/analysis/CP-07b/report.json')
    assert cp08['manuscript_sha256']==digest(source)
    for name,value in cp08['input_sha256'].items(): assert digest(ROOT/name)==value,name
    layout=profile(); author=json.loads((ROOT/'paper/author.json').read_text())
    figures=output.parent/'word_figures'; generate(ROOT/'results/analysis/CP-07b',figures)
    doc=Document(); normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1; normal.paragraph_format.space_after=Pt(4)
    for section in doc.sections:
        values=layout['sections'][0]
        for field,key in [('page_width','width'),('page_height','height'),('top_margin','top'),
                          ('bottom_margin','bottom'),('left_margin','left'),('right_margin','right'),
                          ('header_distance','header'),('footer_distance','footer')]: setattr(section,field,Twips(values[key]))
    ledger=[]; table_index=0; equation_index=0; figure_names=[]
    def paragraph(text='',size=11,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,indent=True):
        p=doc.add_paragraph(); p.alignment=align
        p.paragraph_format.first_line_indent=Inches(.25) if indent else Pt(0)
        run=p.add_run(text); run.font.name='Times New Roman'; run.font.size=Pt(size); run.bold=bold; run.italic=italic
        ledger.append(dict(kind='paragraph',text=text)); return p
    lines=source.read_text(encoding='utf-8').splitlines(); i=0; subsection=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if i==0:
            p=paragraph(line[2:],14,True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
            p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(8)
            paragraph(author['name'],11,True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False).paragraph_format.space_after=Pt(1)
            paragraph(author['affiliation'],9,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
            i+=1; continue
        if line in ('**'+author['name']+'**',author['affiliation']): i+=1; continue
        if line=='## Abstract':
            i+=1
            while not lines[i].strip(): i+=1
            p=paragraph('Abstract—'+lines[i].strip()); p.runs[0].text=lines[i].strip()
            # Insert the reference's bold/italic lead-in without changing content.
            lead=OxmlElement('w:r'); props=OxmlElement('w:rPr'); props.append(OxmlElement('w:b')); props.append(OxmlElement('w:i'))
            lead.append(props); text=OxmlElement('w:t'); text.text='Abstract—'; lead.append(text); p._p.insert(1,lead)
            i+=1; continue
        if line.startswith('**Index terms:'):
            text='Index Terms—'+editorial(line.split(':',1)[1]); paragraph(text)
            section=doc.add_section(WD_SECTION_START.CONTINUOUS)
            cols=section._sectPr.find(qn('w:cols')); cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'440')
            i+=1; continue
        heading=re.match(r'## (\d+)\. (.+)',line)
        if heading:
            subsection=0
            p=paragraph(roman(int(heading[1]))+'. '+heading[2],11,True,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
            p.paragraph_format.space_before=Pt(10); p.paragraph_format.keep_with_next=True; i+=1; continue
        heading=re.match(r'### \d+\.\d+ (.+)',line)
        if heading:
            p=paragraph(chr(65+subsection)+'. '+heading[1],11,italic=True,align=WD_ALIGN_PARAGRAPH.LEFT,indent=False)
            p.paragraph_format.keep_with_next=True; subsection+=1; i+=1; continue
        if line=='## References':
            paragraph('References',11,True,align=WD_ALIGN_PARAGRAPH.LEFT,indent=False).paragraph_format.keep_with_next=True
            i+=1; continue
        if line.startswith('$$'):
            equation_index+=1; p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0)
            equation(p,equation_index); ledger.append(dict(kind='equation',number=equation_index,latex=line.strip('$')))
            i+=1; continue
        picture=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if picture:
            name=Path(picture[2]).stem; figure_names.append(name)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True; p.add_run().add_picture(str(figures/f'{name}.png'),width=Inches(3.3))
            ledger.append(dict(kind='figure',name=name,sha256=digest(figures/f'{name}.png')))
            paragraph(editorial(picture[1]),9,italic=True,indent=False); i+=1; continue
        if line.startswith('**Table '):
            # Number/title are emitted with the table, allowing the related-work table first.
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v) for v in cells): rows.append(cells)
                i+=1
            table_index+=1
            captions=re.findall(r'^\*\*Table \d+\. (.+?)\*\*$',source.read_text(encoding='utf-8'),re.MULTILINE)
            title='Representative Prior Work and Its Difference from This Study' if table_index==1 else editorial(captions[table_index-2])
            for text,bold,italic in [('TABLE '+roman(table_index),True,False),(title,False,True)]:
                p=paragraph(text,9,bold,italic,WD_ALIGN_PARAGRAPH.CENTER,False); p.paragraph_format.keep_with_next=True
            table=doc.add_table(rows=0,cols=len(rows[0])); table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.CENTER
            weights=[max(9,min(24,max(len(row[c]) for row in rows))) for c in range(len(rows[0]))]
            widths=[round(4869*w/sum(weights)) for w in weights]; widths[-1]+=4869-sum(widths)
            for c,width in zip(table.columns,widths): c.width=Twips(width)
            table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),'4869')
            table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                element=OxmlElement('w:'+side); element.set(qn('w:val'),'single'); element.set(qn('w:sz'),'4'); borders.append(element)
            table._tbl.tblPr.append(borders)
            for row_index,values in enumerate(rows):
                row=table.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if row_index==0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell,value,width in zip(row.cells,values,widths):
                    cell.width=Twips(width); p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.line_spacing=1
                    r=p.add_run(value.replace('_','_\u200b')); r.font.name='Times New Roman'; r.font.size=Pt(8); r.bold=row_index==0
            ledger.append(dict(kind='table',number=table_index,rows=rows,title=title))
            continue
        content=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![','$$')):
            content.append(lines[i].strip()); i+=1
        text=editorial(' '.join(content))
        if text:
            is_reference=bool(re.match(r'^\[\d+\]',text))
            p=paragraph(text,9 if is_reference else 11,indent=not is_reference)
            links=list(re.finditer(r'\[([^]]+)\]\(([^)]+)\)',' '.join(content)))
            if is_reference and links:
                label,url=links[-1].groups()
                assert text.endswith(label+'.')
                p.runs[0].text=text[:-len(label)-1]
                hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
                run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
                color=OxmlElement('w:color'); color.set(qn('w:val'),'254B87'); props.append(color)
                size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); props.append(size)
                run.append(props); value=OxmlElement('w:t'); value.text=label; run.append(value)
                hyperlink.append(run); p._p.append(hyperlink); p.add_run('.')
    assert table_index==10 and equation_index==2 and len(figure_names)==7
    doc.core_properties.author=author['name']; doc.core_properties.title=lines[0][2:]
    doc.core_properties.subject='Audited classical search and DQN policy transfer study'
    doc.save(output)
    (output.parent/'word_reference_profile.json').write_text(json.dumps(layout,indent=2)+'\n')
    proof=dict(date='2026-10-07',source_sha256=digest(source),docx_sha256=digest(output),
        reference_sha256=digest(REFERENCE),analysis_sha256=audit['report_sha256'],
        builder_sha256=digest(Path(__file__)),figure_generator_sha256=digest(ROOT/'paper/word_figures.py'),
        author=author,formatting_profile=layout,blocks=ledger,editorial_omissions=OMISSIONS,
        status='DOCX built; structural and native-render checks pending',
        figure_inputs={p.name:digest(p) for p in figures.glob('*') if p.is_file()})
    (output.parent/'word_provenance.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Created {output}; 10 tables, 7 figures, 2 editable numbered equations')
    return output


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'paper/Graph_Search_Research_Paper.docx')
    args=parser.parse_args(); build(output=args.output)
