"""Checks the delivered manuscript and detects rehashed numerical tampering."""
import json
import tempfile
import unittest
from pathlib import Path
from docx import Document
from evaluation.final_protocol import digest
from evaluation.team_protocol import ROOT
from paper.check_word import check
from paper.preview_word import render
import pdfplumber
import shutil


@unittest.skipUnless((ROOT/'paper/Graph_Search_Research_Paper.docx').exists(),
                     'Generate the CP-09 Word deliverable before artifact checks')
class WordManuscriptTests(unittest.TestCase):
    def test_delivered_word_matches_audited_tables_and_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            result=check(ROOT/'paper/Graph_Search_Research_Paper.docx',
                         ROOT/'paper/word_provenance.json',Path(temporary)/'proof.json')
            self.assertTrue(result['all_measured_cells_match'])
            self.assertEqual(result['tables_checked'],10)
            self.assertEqual(result['editable_numbered_equations'],2)

    def test_changed_measured_cell_is_rejected_even_with_new_package_hash(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory=Path(temporary); source=directory/'altered.docx'
            document=Document(ROOT/'paper/Graph_Search_Research_Paper.docx')
            # Result table II, not the literature table: a measured count is corrupted.
            document.tables[1].cell(1,1).text='999/720'; document.save(source)
            provenance=json.loads((ROOT/'paper/word_provenance.json').read_text(encoding='utf-8'))
            provenance['docx_sha256']=digest(source)
            proof=directory/'provenance.json'; proof.write_text(json.dumps(provenance),encoding='utf-8')
            with self.assertRaisesRegex(AssertionError,'Measured table changed'):
                check(source,proof,directory/'must_not_exist.json')
            self.assertFalse((directory/'must_not_exist.json').exists())

    def test_preview_equations_and_figures_stay_inside_pages(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory=Path(temporary); source=directory/'manuscript.docx'
            shutil.copy2(ROOT/'paper/Graph_Search_Research_Paper.docx',source)
            pdf=render(source,directory/'preview.pdf')
            with pdfplumber.open(pdf) as document:
                text='\n'.join(page.extract_text() or '' for page in document.pages)
                self.assertIn('(1)',text); self.assertIn('(2)',text)
                images=[(page,image) for page in document.pages for image in page.images]
                self.assertEqual(len(images),9)  # Seven charts and two equation rasters.
                self.assertTrue(all(im['x0']>=0 and im['x1']<=page.width+.5 and
                    im['top']>=0 and im['bottom']<=page.height+.5 for page,im in images))


if __name__=='__main__': unittest.main()
