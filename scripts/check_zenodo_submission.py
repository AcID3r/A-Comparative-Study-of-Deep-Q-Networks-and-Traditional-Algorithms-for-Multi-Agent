"""Check the author-only preprint revision against the audited CP-09 manuscript."""
import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pdfplumber
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT/'final submission/zenodo'
NAME = 'Madhav_Joshi_Graph_Search_Preprint_v1.0'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value):
    return ' '.join(value.replace('\u200b', '').split())


def check():
    source, pdf_path = FOLDER/(NAME+'.docx'), FOLDER/(NAME+'.pdf')
    accepted_path = ROOT/'paper/Graph_Search_Research_Paper.docx'
    accepted = Document(accepted_path)
    current = Document(source)
    header = ['Madhav Joshi', 'Artificial Intelligence & Machine Learning',
              'Marwadi University', 'madhav.joshi132389@marwadiuniversity.ac.in']
    assert [p.text for p in current.paragraphs[1:5]] == header
    assert current.paragraphs[1].runs[0].bold and current.paragraphs[1].runs[0].font.size.pt == 11
    assert current.core_properties.author == header[0]
    expected = [normalize(p.text) for p in accepted.paragraphs if p.text.strip()]
    actual = [normalize(p.text) for p in current.paragraphs if p.text.strip() and p.text not in (header[1], header[3])]
    assert actual == expected, 'Scientific paragraph content changed'
    cells = lambda d: [[[normalize(c.text) for c in row.cells] for row in table.rows] for table in d.tables]
    assert cells(current) == cells(accepted), 'Measured or literature table changed'
    def package(path):
        with ZipFile(path) as archive:
            images = sorted(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/'))
            xml = etree.fromstring(archive.read('word/document.xml'))
            equations = [normalize(''.join(element.itertext())) for element in xml.findall('.//'+qn('m:oMath'))]
            links = len(xml.findall('.//'+qn('w:hyperlink')))
            return images, equations, links
    assert package(source) == package(accepted_path)
    images, equations, links = package(source)
    assert len(images) == 7 and len(equations) == 2 and links == 9
    provenance = json.loads((ROOT/'results/validation/zenodo-preprint-provenance.json').read_text())
    assert provenance['source_sha256'] == digest(source) and provenance['pdf_sha256'] == digest(pdf_path)
    text = []; pages = []; image_count = 0
    with pdfplumber.open(pdf_path) as document:
        for number, page in enumerate(document.pages, 1):
            content = page.extract_text() or ''
            assert content.strip(), (number, 'empty page')
            assert all(c['x0']>=-.5 and c['x1']<=page.width+.5 and c['top']>=-.5 and c['bottom']<=page.height+.5 for c in page.chars), (number, 'text bounds')
            assert all(im['x0']>=-.5 and im['x1']<=page.width+.5 and im['top']>=-.5 and im['bottom']<=page.height+.5 for im in page.images), (number, 'image bounds')
            text.append(content); image_count += len(page.images)
            pages.append(dict(page=number,characters=len(page.chars),images=len(page.images)))
    joined = '\n'.join(text)
    for value in header + ['89.17%', '24.17%', '22.19%', '123.251', '(1)', '(2)']:
        assert value in joined, ('PDF required content', value)
    for title in ('Introduction', 'Related Work', 'Problem Formulation', 'Methodology', 'Experimental Setup', 'Results', 'Discussion', 'Limitations', 'Conclusion', 'References'):
        assert title in joined, title
    assert image_count == 9, 'Seven charts and two equation rasters must be present'
    assert 'Independent DOCX preview' not in joined
    assert not re.search(r'\bTODO\b|\bTBD\b|\[insert|\[result\]', joined)
    destination = ROOT/'results/validation/zenodo-pages'/digest(pdf_path)[:12]
    destination.mkdir(parents=True,exist_ok=True)
    document = pdfium.PdfDocument(str(pdf_path))
    for number in range(len(document)):
        page = document[number]; bitmap = page.render(scale=1.4)
        bitmap.to_pil().save(destination/f'page_{number+1:02d}.png')
        bitmap.close(); page.close()
    document.close()
    proof = dict(date='2026-10-07', docx_sha256=digest(source), pdf_sha256=digest(pdf_path),
                 audited_scientific_text_unchanged=True, all_ten_tables_unchanged=True,
                 seven_figure_hashes_unchanged=True, equations_and_reference_links_unchanged=True,
                 user_supplied_author_details=header, pdf_renderer=provenance['renderer'],
                 native_word_pagination_verified=False, pages=pages, page_count=len(pages),
                 text_outside_page=0, images_outside_page=0,
                 review_directory=destination.relative_to(ROOT).as_posix(), visual_review='pending inspection')
    (ROOT/'results/validation/zenodo-submission-checks.json').write_text(json.dumps(proof,indent=2)+'\n')
    print(json.dumps(proof,indent=2)); return proof


if __name__ == '__main__':
    check()
